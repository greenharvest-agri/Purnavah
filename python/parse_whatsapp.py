"""
Purnavah WhatsApp Order Parser (catalog-driven)

Paste a WhatsApp order message (from the index.html order form) and this
script extracts the line items + delivery details and appends a row to
orders.xlsx (created from orders_template.xlsx if it doesn't exist yet).

Usage (run from inside python/):
  python parse_whatsapp.py
    -> then paste the message, press Ctrl+Z then Enter (Windows) to finish.

  echo "<message>" | python parse_whatsapp.py   -> non-interactive / scripted use
"""
import sys, os, re
from datetime import date
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, os.path.dirname(__file__))
from load_catalog import build_wa_product_map, load_catalog

# Build maps dynamically from catalog
PRODUCT_MAP = build_wa_product_map()
ALL_COLUMNS = [p["product_name"] for p in load_catalog(active_only=False)]

META_COLS = ["Customer Name", "Phone", "Pincode", "Address", "Notes", "Order Date"]

ORDER_LINE_RE = re.compile(r'•\s*(.+?)\s*×\s*(\d+)\s*=', re.UNICODE)
FIELD_RE      = re.compile(r'^\s*(Name|Phone|Pincode|Address|Notes)\s*:\s*(.*)$', re.IGNORECASE)


def parse_message(text):
    """
    Returns (items, delivery) where:
      items    = list of {"wa_key": str, "product_name": str|None, "qty": int}
      delivery = dict with keys name, phone, pincode, address, notes
    """
    items = []
    for m in ORDER_LINE_RE.finditer(text):
        wa_key = m.group(1).strip()
        qty    = int(m.group(2))
        items.append({
            "wa_key":       wa_key,
            "product_name": PRODUCT_MAP.get(wa_key),
            "qty":          qty,
        })

    delivery = {"name": "", "phone": "", "pincode": "", "address": "", "notes": ""}
    for line in text.splitlines():
        m = FIELD_RE.match(line)
        if not m:
            continue
        key, val = m.group(1).lower(), m.group(2).strip()
        delivery[key] = val

    return items, delivery


def append_to_orders(items, delivery, orders_path="orders.xlsx"):
    """Create orders.xlsx from orders_template.xlsx if needed, then append one row."""
    import openpyxl

    if os.path.exists(orders_path):
        wb = openpyxl.load_workbook(orders_path)
        ws = wb["Orders"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    else:
        template = "orders_template.xlsx" if os.path.exists("orders_template.xlsx") else None
        if template:
            import shutil
            shutil.copy(template, orders_path)
            wb = openpyxl.load_workbook(orders_path)
            ws = wb["Orders"]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Orders"
            headers = META_COLS + ALL_COLUMNS
            ws.append(headers)

    col = {h: i + 1 for i, h in enumerate(headers)}
    row_idx = ws.max_row + 1

    ws.cell(row_idx, col.get("Customer Name", 1), delivery.get("name", ""))
    if "Phone" in col:      ws.cell(row_idx, col["Phone"],      delivery.get("phone", ""))
    if "Pincode" in col:    ws.cell(row_idx, col["Pincode"],    delivery.get("pincode", ""))
    if "Address" in col:    ws.cell(row_idx, col["Address"],    delivery.get("address", ""))
    if "Notes" in col:      ws.cell(row_idx, col["Notes"],      delivery.get("notes", ""))
    if "Order Date" in col: ws.cell(row_idx, col["Order Date"], date.today().isoformat())

    for it in items:
        pname = it["product_name"]
        if pname and pname in col:
            ws.cell(row_idx, col[pname], it["qty"])

    wb.save(orders_path)
    return orders_path, row_idx


def main():
    print("Paste the WhatsApp order message below.")
    print("When done, press Ctrl+Z then Enter (Windows) to finish input.\n")
    text = sys.stdin.read()

    items, delivery = parse_message(text)

    if not items:
        print("No order lines found (expected lines like '• Product Name (1 L) × 2 = ₹880').")
        return

    print("\nParsed order lines:")
    unmatched = 0
    for it in items:
        if it["product_name"]:
            print(f"  {it['wa_key']} → {it['product_name']}  (qty: {it['qty']})")
        else:
            unmatched += 1
            print(f"  {it['wa_key']} → NOT FOUND IN CATALOG  (qty: {it['qty']})")

    print("\nDelivery details:")
    for k in ("name", "phone", "pincode", "address", "notes"):
        print(f"  {k.capitalize()}: {delivery.get(k) or '(blank)'}")

    if unmatched:
        print(f"\nWARNING: {unmatched} item(s) did not match any product in catalog.xlsx — "
              f"they will be skipped when writing to orders.xlsx.")

    if not delivery.get("name"):
        print("\nWARNING: No customer name found — skipping write to orders.xlsx.")
        return

    path, row_idx = append_to_orders(items, delivery)
    print(f"\n✓ Added to {path} (row {row_idx})")


if __name__ == "__main__":
    main()
